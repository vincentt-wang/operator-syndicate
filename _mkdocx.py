# -*- coding: utf-8 -*-
"""從 Operator Insider 文章 HTML 產出真正的 .docx，直接上傳 Google Drive 就能開。

用法: python3 _mkdocx.py <文章html> [輸出目錄] [額外說明檔.txt]

為什麼不是 .html：Vincent 2026-09-15「google drive 沒有的話就給我 doc 文件我上傳雲端」。
拖 HTML 進 Drive 雖然也會轉，但多一步而且格式常跑掉，.docx 是 Drive 原生支援的。
"""
import io,re,html,sys,os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import importlib.util
spec=importlib.util.spec_from_file_location('mk', os.path.join(os.path.dirname(os.path.abspath(__file__)),'_mkdoc.py'))
mk=importlib.util.module_from_spec(spec); spec.loader.exec_module(mk)

TAG = re.compile(r'''<([a-zA-Z][\w-]*)((?:[^>"']|"[^"]*"|'[^']*')*)>''')

def blocks_of(src):
    t=io.open(src,encoding='utf-8').read()
    art=t[t.index('<article>'):t.index('</article>')]
    out=[]
    for m in TAG.finditer(art):
        tag,attr=m.group(1).lower(),m.group(2)
        if 'data-en="' not in attr: continue
        kind=mk.classify(tag,attr)
        if not kind: continue
        en=html.unescape(re.sub(r'<[^>]+>','',re.search(r'data-en="([^"]*)"',attr).group(1)))
        zh=re.search(r'data-zh="([^"]*)"',attr)
        zh=html.unescape(re.sub(r'<[^>]+>','',zh.group(1))) if zh else ''
        if not (en.strip() or zh.strip()): continue
        out.append((kind,en,zh))
    meta={}
    w=re.search(r'class="opi-au-n">([^<]+)<',t); meta['who']=w.group(1).strip() if w else ''
    r=re.search(r'class="opi-au-t"[^>]*data-zh="([^"]*)"',t); meta['role']=html.unescape(r.group(1)) if r else ''
    h=re.search(r'<h1[^>]*data-zh="([^"]*)"',t); meta['h1zh']=html.unescape(h.group(1)) if h else ''
    h=re.search(r'<h1[^>]*data-en="([^"]*)"',t); meta['h1en']=html.unescape(h.group(1)) if h else ''
    return out, meta

def add(doc, text, style=None, italic=False, bold=False, grey=False, indent=0):
    p=doc.add_paragraph(style=style)
    if indent: p.paragraph_format.left_indent=Inches(indent)
    r=p.add_run(text); r.italic=italic; r.bold=bold
    if grey: r.font.color.rgb=RGBColor(0x66,0x66,0x66)
    return p

def build(src, outdir=None, notes_file=None):
    blocks, meta = blocks_of(src)
    doc=Document()
    st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(11)

    doc.add_heading('Operator Insider · '+meta['h1zh'], 0)
    add(doc, '%s · %s' % (meta['who'], meta['role']), bold=True)
    for s in ['這一份是給作者改字用的。情境圖、分析圖、互動元素與動畫都已經做好了，在網頁上，不在這份文件裡。標「[圖／互動]」的地方就是它們的位置，下面那行「圖內文字」就是圖上會出現的字，要改直接改那一行。',
              '中文與英文是兩個平行版本，不是互相翻譯，兩邊都可以各自改。',
              '所有數字後面都附了來源與發布日，在最後一段。如果覺得哪一個數字不該用，直接刪掉那一句就好，不用替我們找替代。']:
        add(doc, s, italic=True, grey=True)

    if notes_file and os.path.exists(notes_file):
        doc.add_page_break()
        for line in io.open(notes_file,encoding='utf-8').read().split('\n'):
            line=line.strip()
            if not line: continue
            if line.startswith('#'): add(doc, line.lstrip('# ').strip(), style='Heading 2')
            else: add(doc, line, italic=True)

    # 每個 section 後面留一個「你的修改」空欄。
    # Vincent 2026-09-15：「每一段、每個 section，它可以分開改⋯⋯不改的話就留空白」。
    def slot(doc, n, lang):
        lab = ('第 %d 段　你的修改' % n) if lang=='zh' else ('Section %d — your edits' % n)
        hint = ('這一段不用改就整段留白。要改的話直接在這裡寫，不必動上面的原文。'
                if lang=='zh' else
                'Leave this blank if the section is fine. Write your version here; no need to touch the text above.')
        p1=doc.add_paragraph(); r=p1.add_run('▸ '+lab); r.bold=True
        r.font.color.rgb=RGBColor(0xC2,0x41,0x0C)
        add(doc, hint, italic=True, grey=True)
        doc.add_paragraph()   # 留白讓人寫
        doc.add_paragraph('_'*58)

    for lang,title in (('zh','中文版'),('en','English version')):
        doc.add_page_break()
        doc.add_heading(title, 1)
        pend=[]; secn=0; started=False
        def flush():
            if pend:
                add(doc, '圖內文字：'+'　｜　'.join(pend), italic=True, grey=True, indent=0.3)
                pend.clear()
        for k,en,zh in blocks:
            s=(zh if lang=='zh' else en).strip()
            if not s: continue
            if k=='label': pend.append(s); continue
            flush()
            if k in ('h2','h3'):
                if started: slot(doc, secn, lang)
                secn+=1; started=True
                add(doc, s, style='Heading 2')
            elif k=='h1': add(doc, s, style='Heading 1')
            elif k=='fig': add(doc, '[圖／互動] '+s, bold=True)
            elif k=='quote': add(doc, s, italic=True, indent=0.4)
            elif k=='note': add(doc, s, italic=True, grey=True)
            elif k=='ref': add(doc, s, grey=True)
            else:
                if not started: secn+=1; started=True
                add(doc, s)
        flush()
        if started: slot(doc, secn, lang)

    outdir=outdir or os.path.dirname(src)
    base=os.path.splitext(os.path.basename(src))[0]
    out=os.path.join(outdir,'%s.docx'%base)
    doc.save(out)
    return out, len(blocks)

if __name__=='__main__':
    o,n=build(sys.argv[1], sys.argv[2] if len(sys.argv)>2 else None,
              sys.argv[3] if len(sys.argv)>3 else None)
    print(o); print('  區塊 %d'%n)
